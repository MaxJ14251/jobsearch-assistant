"""`jsa import-resume` (Plan 8): a resume becomes a DRAFT profile.

Every fixture here is a fictional person built in the test with python-docx.
The promises under test: identity never reaches the prompt, nothing reworded
survives into the draft, the credential is never filled in, and the live
profile is never written.
"""

import io
import json
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
from unittest import mock

import yaml
from docx import Document

from jsa import cli, db, doctor, llm
from jsa import resume_import as ri
from jsa.llm import Usage

ROOT = Path(__file__).resolve().parents[1]

NAME = "Jordan Avery Quill"
EMAIL = "jordan.quill@example.com"
PHONE = "(555) 010-2244"
STREET = "48 Larkspur Lane"
ZIP = "97024"
LINKEDIN = "linkedin.com/in/jordan-quill-example"

KEPT = "Resolved customer tickets for a 40-person warehouse team within one business day."
ORIGINAL = "Helped lead the move of 300 product records into a new inventory system."
REWORDED = "Led the move of 300 product records into a new inventory system."


def build_resume(path: Path, *, header=True, name=True) -> Path:
    doc = Document()
    contact = f"{EMAIL} | {PHONE} | {STREET}, Fairview, OR {ZIP} | {LINKEDIN}"
    if header:
        hdr = doc.sections[0].header
        hdr.paragraphs[0].text = NAME if name else "Resume"
        hdr.add_paragraph(contact)
    else:
        if name:
            doc.add_paragraph(NAME)
        doc.add_paragraph(contact)
    doc.add_paragraph("EXPERIENCE")
    table = doc.add_table(rows=1, cols=2)
    left, right = table.rows[0].cells
    left.paragraphs[0].text = "Northwind Supply"
    left.add_paragraph("Support Specialist")
    left.add_paragraph("2021 - Present")
    right.paragraphs[0].text = "• " + KEPT
    right.add_paragraph("• " + ORIGINAL)
    doc.add_paragraph("EDUCATION")
    doc.add_paragraph("Lakeside Community College, Business coursework, 2018 - 2020")
    doc.add_paragraph("SKILLS")
    doc.add_paragraph("Python, SQL, Zendesk")
    doc.save(str(path))
    return path


def model_reply():
    """What a model might return, including the two kinds of mistake."""
    return {
        "experience": [
            {"company": "Northwind Supply", "title": "Support Specialist",
             "location": None, "start": "2021", "end": None, "current": True,
             "family": "technical_field",
             "bullets": [{"text": KEPT, "tags": ["Customer Facing", "support"]},
                         {"text": REWORDED, "tags": ["migration"]}]},
            {"company": "Globex Corporation", "title": "Analyst",
             "start": "2019", "end": "2020",
             "bullets": [{"text": "Built dashboards.", "tags": []}]},
        ],
        "projects": [],
        "certifications": [],
        "education": [{"institution": "Lakeside Community College",
                       "credential": "BS Business", "field": "Business",
                       "start": "2018", "end": "2020",
                       "line": "Lakeside Community College, Business coursework, 2018 - 2020"}],
        "skills": ["Python", "SQL", "Kubernetes"],
        "target_titles": ["Support Engineer", "Technical Support Specialist"],
    }


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.resume = build_resume(self.tmp / "resume.docx")


class TestReading(Base):
    def test_header_and_table_cells_are_read(self):
        lines = ri.read_docx(self.resume)
        self.assertEqual(lines[0], NAME)
        self.assertIn("Northwind Supply", lines)
        self.assertIn(KEPT, lines)  # the bullet glyph is stripped
        self.assertIn("Python, SQL, Zendesk", lines)

    def test_a_pdf_is_refused_with_the_fix(self):
        pdf = self.tmp / "resume.pdf"
        pdf.write_bytes(b"%PDF-1.4")
        with self.assertRaises(ri.ResumeReadError) as ctx:
            ri.read_docx(pdf)
        self.assertIn("Save your resume as .docx", str(ctx.exception))


class TestIdentity(Base):
    def test_contact_details_are_found_and_removed(self):
        ident, redacted = ri.split_identity(ri.read_docx(self.resume))
        self.assertEqual(ident.full_name, NAME)
        self.assertEqual(ident.email, EMAIL)
        self.assertEqual(ident.phone, PHONE)
        self.assertEqual(ident.postal_code, ZIP)
        self.assertEqual((ident.city, ident.state), ("Fairview", "OR"))
        self.assertEqual(ident.street, STREET)
        self.assertTrue(ident.links["linkedin"].endswith(LINKEDIN))
        text = "\n".join(redacted)
        for value in (NAME, EMAIL, PHONE, STREET, ZIP, LINKEDIN, "Fairview"):
            self.assertNotIn(value, text)
        self.assertIn(KEPT, text)

    def test_us_phone_formats(self):
        for phone in ("555-010-2244", "555.010.2244", "(555) 010-2244",
                      "+1 555 010 2244", "5550102244"):
            ident, redacted = ri.split_identity([f"Call {phone}"])
            self.assertEqual(ident.phone, phone)
            self.assertNotIn(phone, redacted[0])

    def test_digits_in_a_link_are_not_a_zip(self):
        ident, redacted = ri.split_identity(
            ["Pat Example", "github.com/patexample24680 | Fairview, OR"])
        self.assertIsNone(ident.postal_code)
        self.assertEqual(ident.links["github"], "https://github.com/patexample24680")

    def test_contact_block_in_the_body_works_too(self):
        path = build_resume(self.tmp / "body.docx", header=False)
        ident, _ = ri.split_identity(ri.read_docx(path))
        self.assertEqual(ident.full_name, NAME)

    def test_no_name_line_leaves_the_name_unset(self):
        path = build_resume(self.tmp / "noname.docx", name=False)
        ident, _ = ri.split_identity(ri.read_docx(path))
        self.assertIsNone(ident.full_name)
        draft = ri.render_draft(ident, ri.Verified(), "noname.docx")
        self.assertIn("full_name: null  # TODO", draft)


class TestExtractAndVerify(Base):
    def test_the_prompt_carries_no_identity_value(self):
        ident, redacted = ri.split_identity(ri.read_docx(self.resume))
        seen = {}

        def capture(prompt, **kw):
            seen["prompt"] = prompt + (kw.get("system") or "")
            return model_reply(), Usage()

        with mock.patch("jsa.resume_import.llm.complete_json", side_effect=capture):
            ri.extract(redacted, ident)
        for value in (NAME, EMAIL, PHONE, "0102244", STREET, ZIP, LINKEDIN):
            self.assertNotIn(value.lower(), seen["prompt"].lower())
        self.assertIn(KEPT, seen["prompt"])

    def test_a_leak_stops_before_the_network(self):
        ident, redacted = ri.split_identity(ri.read_docx(self.resume))
        redacted.append(f"Reach me at {EMAIL}")  # a line the redaction missed
        from jsa.tailor import IdentityLeakError
        with mock.patch("jsa.resume_import.llm.complete_json") as call, \
             self.assertRaises(IdentityLeakError):
            ri.extract(redacted, ident)
        call.assert_not_called()

    def verified(self):
        ident, redacted = ri.split_identity(ri.read_docx(self.resume))
        return ri.verify(model_reply(), "\n".join(redacted))

    def test_a_reworded_bullet_is_dropped_and_reported(self):
        v = self.verified()
        texts = [b["text"] for b in v.experience[0]["bullets"]]
        self.assertEqual(texts, [KEPT])
        self.assertTrue(any(d.what == REWORDED and d.why == ri.NOT_VERBATIM
                            for d in v.dropped))

    def test_an_invented_employer_is_dropped(self):
        v = self.verified()
        self.assertEqual([e["company"] for e in v.experience], ["Northwind Supply"])
        self.assertTrue(any(d.what == "Globex Corporation" for d in v.dropped))

    def test_a_skill_not_in_the_resume_is_dropped(self):
        self.assertEqual(self.verified().skills, ["Python", "SQL"])

    def test_the_credential_is_never_filled_in(self):
        v = self.verified()
        self.assertNotIn("credential", v.education[0])
        draft = yaml.safe_load(ri.render_draft(ri.Identity(), v, "r.docx"))
        self.assertIsNone(draft["education"][0]["credential"])
        self.assertNotIn("BS Business", ri.render_draft(ri.Identity(), v, "r.docx"))

    def test_a_heading_keeps_its_name_when_its_link_was_removed(self):
        # Found on the first real run: "Project Name - <repo link>" came back
        # from the model as "Project Name - [contact]" and failed the check.
        reply = {"projects": [{"name": "Northwind Supply — [contact]",
                               "bullets": [{"text": KEPT}]}]}
        v = ri.verify(reply, "Northwind Supply — [contact]\n" + KEPT)
        self.assertEqual(v.projects[0]["name"], "Northwind Supply")
        self.assertEqual(len(v.projects[0]["bullets"]), 1)

    def test_a_dropped_project_lists_its_bullets_too(self):
        reply = {"projects": [{"name": "Invented Project",
                               "bullets": [{"text": KEPT}]}]}
        v = ri.verify(reply, KEPT)
        self.assertEqual(v.projects, [])
        self.assertTrue(any(d.what == KEPT and "with its project" in d.why
                            for d in v.dropped))

    def test_tags_are_normalized_suggestions(self):
        self.assertEqual(self.verified().experience[0]["bullets"][0]["tags"],
                         ["customer-facing", "support"])

    def test_a_date_with_a_year_not_in_the_resume_is_dropped(self):
        reply = model_reply()
        reply["experience"][0]["start"] = "2016-04"
        ident, redacted = ri.split_identity(ri.read_docx(self.resume))
        v = ri.verify(reply, "\n".join(redacted))
        self.assertIsNone(v.experience[0]["start"])


class TestCommand(Base):
    def setUp(self):
        super().setUp()
        self.profile_dir = self.tmp / "profile"
        self.profile_dir.mkdir()
        self.live = self.profile_dir / "master_profile.yaml"
        self.live.write_bytes(b"identity: {full_name: Someone Else}\n")
        self.draft = self.profile_dir / "master_profile.draft.yaml"
        self.dbfile = self.tmp / "t.db"
        db.init_db(self.dbfile)
        con = db.connect(self.dbfile)
        con.execute("INSERT INTO companies (id,name,slug) VALUES (1,'Acme','acme')")
        con.execute(
            "INSERT INTO jobs (id,company_id,title,url,description,location,remote) "
            "VALUES (5,1,'Support Engineer','https://acme.test/5',"
            "'Python and SQL support for customers.','Remote','remote')")
        con.execute(
            "INSERT INTO jobs (id,company_id,title,url,description,location,remote) "
            "VALUES (6,1,'Account Executive','https://acme.test/6',"
            "'Quota-carrying sales.','Remote','remote')")
        con.commit()
        con.close()
        real_connect = db.connect
        for patch in (
            mock.patch("jsa.config.PROFILE_PATH", self.live),
            mock.patch("jsa.cli.DB_PATH", self.dbfile),
            mock.patch("jsa.db.connect",
                       side_effect=lambda *a, **k: real_connect(self.dbfile)),
            mock.patch("jsa.resume_import.llm.complete_json",
                       return_value=(model_reply(), Usage())),
        ):
            patch.start()
            self.addCleanup(patch.stop)

    def run_cli(self, *extra):
        out, err = io.StringIO(), io.StringIO()
        with redirect_stdout(out), redirect_stderr(err):
            code = cli.main(["import-resume", str(self.resume), *extra])
        return code, out.getvalue(), err.getvalue()

    def test_writes_a_draft_and_never_the_live_profile(self):
        before = self.live.read_bytes()
        code, out, _ = self.run_cli()
        self.assertEqual(code, 0, out)
        self.assertEqual(self.live.read_bytes(), before)
        self.assertTrue(self.draft.exists())
        self.assertIn("imported 1 job(s)", out)
        self.assertIn("dropped", out)
        self.assertIn("#5 Acme", out)
        self.assertIn("jsa doctor", out.splitlines()[-2])

    def test_the_draft_has_the_example_shape(self):
        self.run_cli()
        draft = yaml.safe_load(self.draft.read_text(encoding="utf-8"))
        example = yaml.safe_load(
            (ROOT / "profile" / "master_profile.example.yaml").read_text(encoding="utf-8"))
        self.assertEqual(set(draft), set(example))
        bullet = draft["experience"][0]["bullets"][0]
        self.assertEqual(bullet["id"], "b_northwind_supply_1")
        self.assertEqual(bullet["text"], KEPT)

    def test_doctor_lists_the_undecided_items(self):
        code, out, _ = self.run_cli()
        draft = yaml.safe_load(self.draft.read_text(encoding="utf-8"))
        whats = [f.what for f in doctor.run(draft, None).blocking]
        for expected in ("work_authorization is unanswered",
                         "compensation_floor_usd is undecided",
                         "No summary variants"):
            self.assertIn(expected, whats)
        self.assertTrue(any("credential" in w for w in whats))
        self.assertIn("work_authorization is unanswered", out)

    def test_an_existing_draft_needs_force(self):
        self.draft.write_text("mine\n", encoding="utf-8")
        code, _, err = self.run_cli()
        self.assertEqual(code, 2)
        self.assertIn("--force", err)
        self.assertEqual(self.draft.read_text(encoding="utf-8"), "mine\n")
        code, _, _ = self.run_cli("--force")
        self.assertEqual(code, 0)

    def test_out_cannot_be_the_live_profile(self):
        for target in (self.live, self.tmp / "elsewhere" / "master_profile.yaml"):
            before = self.live.read_bytes()
            code, _, err = self.run_cli("--out", str(target))
            self.assertEqual(code, 2)
            self.assertIn("refused", err)
            self.assertEqual(self.live.read_bytes(), before)

    def test_no_api_key_refuses_cleanly(self):
        with mock.patch("jsa.resume_import.llm.complete_json",
                        side_effect=llm.LLMError("no API key set")):
            code, _, err = self.run_cli()
        self.assertEqual(code, 2)
        self.assertIn("no API key", err)
        self.assertFalse(self.draft.exists())

    def test_preview_writes_nothing(self):
        self.run_cli()
        draft = yaml.safe_load(self.draft.read_text(encoding="utf-8"))
        con = db.connect(self.dbfile)
        self.addCleanup(con.close)
        before = con.total_changes
        top = ri.preview(con, draft)
        self.assertEqual(con.total_changes, before)
        self.assertEqual(top[0].job_id, 5)

    def test_preview_explains_a_draft_it_cannot_score(self):
        draft = {"job_search_preferences": {"target_titles": [], "locations": []}}
        con = db.connect(self.dbfile)
        self.addCleanup(con.close)
        with self.assertRaises(ri.ConfigError):
            ri.preview(con, draft)


class TestKeptOutOfGit(unittest.TestCase):
    def test_draft_and_resume_files_are_ignored(self):
        lines = (ROOT / ".gitignore").read_text(encoding="utf-8").splitlines()
        for pattern in ("profile/*.draft.yaml", "profile/*.docx", "profile/*.pdf"):
            self.assertIn(pattern, lines)

    def test_the_scanner_never_reads_the_draft(self):
        from tools.scan_secrets import NEVER_SCAN
        self.assertIn("master_profile.draft.yaml", NEVER_SCAN)


if __name__ == "__main__":
    unittest.main()
