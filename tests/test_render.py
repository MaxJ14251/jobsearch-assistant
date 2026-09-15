"""Document rendering tests.

The load-bearing one is TestAtsReadability: generate a .docx, then pull the
text back out programmatically. If extraction fails here, an applicant tracking
system fails too — and the candidate never learns why.

This project's own source resume was a ReportLab PDF that needed a hand-written
ASCII85 + Flate decoder to read. That is the failure being designed against.
"""

import json
import sqlite3
import tempfile
import unittest
from pathlib import Path

from jsa import db, render
from jsa.render import extract_text, render_cover_letter, render_resume
from jsa.tailor import DraftBullet, TailoredDraft, collect_bullets
from tests.test_tailor import PROFILE, credential_says_not_conferred

JOB = {
    "id": None,
    "title": "Forward Deployed Engineer",
    "company": "Example Corp",
    "description": "Python, LLM and customer-facing work.",
    "track": "engineering",
}


def sample_draft():
    sources = collect_bullets(PROFILE)
    ids = ["b_vid_design", "b_game_build", "b_riv_sell", "b_inst_install"]
    return TailoredDraft(
        summary=PROFILE["summaries"][0]["text"],
        bullets=[DraftBullet(source_id=i, text=sources[i].text) for i in ids],
        keywords_matched=["Python", "LLM"],
        keywords_missing=["Kubernetes", "PyTorch"],
        model="test-model",
    )


class TestAtsReadability(unittest.TestCase):
    """Generate, then read back. The round trip is the requirement."""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.draft = sample_draft()
        self.path = render_resume(self.draft, PROFILE, JOB, self.tmp / "resume.docx")

    def test_file_exists_and_is_not_empty(self):
        self.assertTrue(self.path.exists())
        self.assertGreater(self.path.stat().st_size, 2000)

    def test_text_extracts(self):
        text = extract_text(self.path)
        self.assertGreater(len(text), 400, "almost nothing came back out")

    def test_every_generated_bullet_survives_the_round_trip(self):
        """A bullet that doesn't extract is a bullet the ATS never sees."""
        text = " ".join(extract_text(self.path).split()).lower()
        for bullet in self.draft.bullets:
            probe = " ".join(bullet.text.split())[:60].lower()
            self.assertIn(probe, text, f"{bullet.source_id} did not extract")

    def test_identity_is_present_because_it_is_merged_locally(self):
        """Never sent to a model — but it must reach the page."""
        text = extract_text(self.path)
        ident = PROFILE["identity"]
        self.assertIn(ident["full_name"], text)
        self.assertIn(ident["email"], text)
        self.assertIn(ident["phone"], text)

    def test_education_field_of_study_reaches_the_page(self):
        """The field is what 'BS in X or equivalent' postings care about."""
        text = extract_text(self.path).lower()
        field = (PROFILE["education"][0].get("field") or "").lower()
        self.assertIn(field, text)

    @unittest.skipUnless(credential_says_not_conferred(PROFILE),
                         "this profile has a conferred credential")
    def test_unconferred_credential_is_reproduced_verbatim(self):
        """The document must never soften what the profile states."""
        text = extract_text(self.path).lower()
        self.assertIn("not conferred", text)

    def test_in_development_projects_are_labelled(self):
        self.assertIn("in development", extract_text(self.path).lower())

    def test_summary_and_headings_present(self):
        text = extract_text(self.path).upper()
        for heading in ("SUMMARY", "EXPERIENCE", "SKILLS", "EDUCATION"):
            self.assertIn(heading, text)

    def test_no_layout_tables_or_images(self):
        """ATS parsers mangle table layouts and cannot read images at all."""
        from docx import Document
        doc = Document(str(self.path))
        self.assertEqual(len(doc.tables), 0, "tables break ATS parsing")
        self.assertEqual(
            len(doc.inline_shapes), 0, "images are invisible to a parser")

    def test_no_content_hidden_in_headers_or_footers(self):
        from docx import Document
        doc = Document(str(self.path))
        for section in doc.sections:
            for part in (section.header, section.footer):
                text = " ".join(p.text for p in part.paragraphs).strip()
                self.assertEqual(text, "", "content in a header/footer is lost")

    def test_bullets_are_real_list_paragraphs(self):
        """Not a literal bullet character the parser has to guess at."""
        from docx import Document
        doc = Document(str(self.path))
        styles = [p.style.name for p in doc.paragraphs]
        self.assertIn("List Bullet", styles)

    def test_cover_letter_also_round_trips(self):
        body = ("I am writing about the Forward Deployed Engineer role.\n\n"
                "I build LLM automation pipelines in Python.\n\n"
                "Thank you for your time.")
        path = render_cover_letter(
            self.draft, PROFILE, JOB, body, self.tmp / "cover.docx")
        text = extract_text(path)
        self.assertIn("Forward Deployed Engineer", text)
        self.assertIn("LLM automation pipelines", text)
        self.assertIn(PROFILE["identity"]["full_name"], text)


class TestProvenance(unittest.TestCase):
    """A bullet_ids array nobody can resolve is decorative, not provenance."""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.dbfile = self.tmp / "t.db"
        db.init_db(self.dbfile)
        self.con = db.connect(self.dbfile)
        self.con.execute(
            "INSERT INTO companies (id, name, slug) VALUES (1,'Example','example')")
        self.con.execute(
            "INSERT INTO jobs (id, company_id, title, url) "
            "VALUES (1,1,'Forward Deployed Engineer','https://example.com/j/1')")
        self.draft = sample_draft()

    def tearDown(self):
        self.con.close()

    def test_round_trip_every_bullet_id_resolves(self):
        doc_id = render.record(
            self.con, job_id=1, kind="resume", path=self.tmp / "r.docx",
            draft=self.draft, prompt_hash="abc123")
        row = self.con.execute(
            "SELECT * FROM documents WHERE id = ?", (doc_id,)).fetchone()
        known = set(collect_bullets(PROFILE))
        stored = json.loads(row["bullet_ids"])
        self.assertTrue(stored)
        for bid in stored:
            self.assertIn(bid, known, f"dangling provenance id {bid!r}")

    def test_keyword_gap_is_recorded_honestly(self):
        doc_id = render.record(
            self.con, job_id=1, kind="resume", path=self.tmp / "r.docx",
            draft=self.draft, prompt_hash="abc123")
        row = self.con.execute(
            "SELECT * FROM documents WHERE id = ?", (doc_id,)).fetchone()
        missing = json.loads(row["keywords_missing"])
        self.assertIn("Kubernetes", missing)
        self.assertIn("PyTorch", missing)

    def test_versions_increment_rather_than_overwrite(self):
        for expected in (1, 2, 3):
            doc_id = render.record(
                self.con, job_id=1, kind="resume", path=self.tmp / "r.docx",
                draft=self.draft, prompt_hash="abc")
            v = self.con.execute(
                "SELECT version FROM documents WHERE id = ?", (doc_id,)
            ).fetchone()[0]
            self.assertEqual(v, expected)

    def test_model_and_prompt_hash_are_stored(self):
        doc_id = render.record(
            self.con, job_id=1, kind="resume", path=self.tmp / "r.docx",
            draft=self.draft, prompt_hash="deadbeef")
        row = self.con.execute(
            "SELECT model, prompt_hash FROM documents WHERE id = ?", (doc_id,)
        ).fetchone()
        self.assertEqual(row["model"], "test-model")
        self.assertEqual(row["prompt_hash"], "deadbeef")


class TestOutputIsGitignored(unittest.TestCase):
    def test_generated_documents_are_never_committed(self):
        from jsa.config import ROOT
        ignored = (ROOT / ".gitignore").read_text(encoding="utf-8")
        self.assertIn("output/", ignored)


if __name__ == "__main__":
    unittest.main()
