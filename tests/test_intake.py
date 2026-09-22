"""Adding one job by hand, and reading a document without downloading it.

n13. The operator found jobs outside the feeds and had no way in, and reading
a draft meant downloading a .docx and opening Word.

Nothing here touches the network: every fetch is patched, and a test fails if
an unsupported link would have been fetched at all.
"""

import shutil
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from jsa import db, intake
from jsa.config import Preferences
from jsa.intake import IntakeError, Link, parse_link
from jsa.sources import FetchResult, content_hash


def prefs() -> Preferences:
    return Preferences.from_profile({"job_search_preferences": {
        "target_titles": ["Support Engineer"],
        "locations": ["Remote (US)"],
        "max_years_experience": 3,
        "compensation_floor_usd": "no_floor",
    }, "ats_keywords": {"have": ["Python"]}})


def posting(external_id="4722512005", title="Support Engineer",
            description="Help customers debug Python. " * 20, location="Remote (US)"):
    return {"external_id": external_id, "title": title, "department": None,
            "location": location, "remote": "remote", "employment_type": "unknown",
            "url": f"https://job-boards.greenhouse.io/acme/jobs/{external_id}",
            "description": description, "description_hash": content_hash(description),
            "posted_at": None}


class TestC_LinksParseToTheirBoard(unittest.TestCase):
    CASES = {
        "https://job-boards.greenhouse.io/scaleai/jobs/4722512005":
            Link("greenhouse", "scaleai", "4722512005"),
        "https://boards.greenhouse.io/discord/jobs/8806482002?gh_src=x":
            Link("greenhouse", "discord", "8806482002"),
        "https://boards.greenhouse.io/embed/job_app?for=acme&token=123":
            Link("greenhouse", "acme", "123"),
        "https://jobs.lever.co/system1/84bab2c5-da41-4eea-bab7-31ddcdfc8cbc/apply":
            Link("lever", "system1", "84bab2c5-da41-4eea-bab7-31ddcdfc8cbc"),
        "https://jobs.ashbyhq.com/replit/A45CA7BA-21F9-464D-8C00-EB5361B3C9F4":
            Link("ashby", "replit", "a45ca7ba-21f9-464d-8c00-eb5361b3c9f4"),
    }

    def test_each_shape(self):
        for url, expected in self.CASES.items():
            with self.subTest(url=url):
                self.assertEqual(parse_link(url), expected)

    def test_workday_with_and_without_a_locale(self):
        for url in (
            "https://servicetitan.wd1.myworkdayjobs.com/ServiceTitan/job/Glendale-CA/Rep_JR114406",
            "https://servicetitan.wd1.myworkdayjobs.com/en-US/ServiceTitan/job/Glendale-CA/Rep_JR114406",
        ):
            with self.subTest(url=url):
                link = parse_link(url)
                self.assertEqual((link.kind, link.tenant, link.wd, link.board, link.job_id),
                                 ("workday", "servicetitan", "wd1", "ServiceTitan", "JR114406"))
                self.assertEqual(link.path, "/job/Glendale-CA/Rep_JR114406")

    def test_look_alikes_and_other_sites_are_not_boards(self):
        for url in (
            "https://boards.greenhouse.io.evil.test/acme/jobs/1",
            "https://evil.test/job-boards.greenhouse.io/acme/jobs/1",
            "https://jobs.lever.co.evil.test/acme/84bab2c5-da41-4eea-bab7-31ddcdfc8cbc",
            "https://www.linkedin.com/jobs/view/4012345678",
            "https://www.indeed.com/viewjob?jk=abc",
            "file:///C:/Users/me/.env",
            "javascript:alert(1)",
            "https://job-boards.greenhouse.io/acme",            # a board, not a posting
            "https://jobs.lever.co/acme/not-a-uuid",
            "",
        ):
            with self.subTest(url=url):
                self.assertIsNone(parse_link(url))


class TrackerCase(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.dir, True)
        db.init_db(self.dir / "t.db")
        self.con = db.connect(self.dir / "t.db")
        self.addCleanup(self.con.close)
        patcher = mock.patch.object(intake, "load_sources", return_value=[])
        patcher.start()
        self.addCleanup(patcher.stop)

    def jobs(self) -> int:
        return self.con.execute("SELECT COUNT(*) FROM jobs").fetchone()[0]


class TestD_AnUnsupportedLinkIsNeverFetched(TrackerCase):
    def test_nothing_is_requested_and_the_way_forward_is_named(self):
        with mock.patch("jsa.sources.fetch") as fetch, \
                mock.patch("jsa.sources._client") as client:
            with self.assertRaises(IntakeError) as ctx:
                intake.add_link(self.con, "https://www.linkedin.com/jobs/view/1", prefs())
        fetch.assert_not_called()
        client.assert_not_called()
        self.assertIn("paste", str(ctx.exception).lower())
        self.assertEqual(self.jobs(), 0)

    def test_the_request_goes_to_the_api_built_from_the_board(self):
        """The link picks which posting. It never picks the host."""
        with mock.patch("jsa.sources.fetch",
                        return_value=FetchResult(True, [posting()], "ok")) as fetch:
            intake.add_link(self.con,
                            "https://job-boards.greenhouse.io/acme/jobs/4722512005", prefs())
        fetch.assert_called_once_with({"kind": "greenhouse", "board": "acme"})

    def test_a_closed_posting_is_reported_not_invented(self):
        with mock.patch("jsa.sources.fetch", return_value=FetchResult(True, [], "ok")):
            with self.assertRaises(IntakeError) as ctx:
                intake.add_link(self.con, "https://job-boards.greenhouse.io/acme/jobs/9", prefs())
        self.assertIn("closed", str(ctx.exception))
        self.assertEqual(self.jobs(), 0)


class TestE_OnePostingOneRow(TrackerCase):
    URL = "https://job-boards.greenhouse.io/acme/jobs/4722512005"

    def add(self, **kw):
        with mock.patch("jsa.sources.fetch",
                        return_value=FetchResult(True, [posting()], "ok")):
            return intake.add_link(self.con, self.URL, prefs(), **kw)

    def test_twice_is_the_same_row(self):
        first, second = self.add(), self.add()
        self.assertTrue(first.new)
        self.assertFalse(second.new)
        self.assertEqual(first.job_id, second.job_id)
        self.assertEqual(self.jobs(), 1)

    def test_naming_it_later_renames_rather_than_duplicating(self):
        self.add()
        self.add(company="Acme Robotics")
        rows = self.con.execute("SELECT name FROM companies").fetchall()
        self.assertEqual([r["name"] for r in rows], ["Acme Robotics"])

    def test_a_guess_never_overwrites_a_real_name(self):
        self.add(company="Acme Robotics")
        added = self.add()
        self.assertEqual(added.company, "Acme Robotics")

    def test_a_job_discovery_already_stored_is_found_not_duplicated(self):
        """Same source name as discovery: <slug>-<kind>."""
        entry = {"company": "Acme Corp", "slug": "acme", "kind": "greenhouse",
                 "board": "acme"}
        company = db.upsert_company(self.con, name="Acme Corp", slug="acme")
        source = db.upsert_source(self.con, name="acme-greenhouse", kind="greenhouse",
                                  url="https://x", company_id=company)
        db.upsert_job(self.con, {**posting(), "company_id": company,
                                 "source_id": source})
        with mock.patch.object(intake, "load_sources", return_value=[entry]):
            added = self.add()
        self.assertFalse(added.new)
        self.assertEqual(self.jobs(), 1)
        self.assertEqual(added.company, "Acme Corp")

    def test_an_existing_application_is_reported(self):
        added = self.add()
        self.con.execute("INSERT INTO applications (job_id, status) VALUES (?, 'applied')",
                         (added.job_id,))
        self.assertEqual(self.add().status, "applied")

    def test_a_posting_the_filters_reject_is_kept_with_the_reason(self):
        senior = posting(description="Requires 10+ years of experience. " * 20)
        with mock.patch("jsa.sources.fetch",
                        return_value=FetchResult(True, [senior], "ok")):
            added = intake.add_link(self.con, self.URL, prefs())
        self.assertEqual(self.jobs(), 1)
        self.assertEqual(added.score, 0)
        self.assertTrue(any("kept because you chose it" in w for w in added.warnings))


class TestF_APastedPostingIsStoredAsPasted(TrackerCase):
    TEXT = ("We need a Support Engineer.\n\nRequirements:\n- Python\n- 2+ years\n"
            "<b>not markup</b> & other characters " + "x" * 200)

    def paste(self, **kw):
        args = {"company": "Acme", "title": "Support Engineer", "text": self.TEXT,
                "prefs": prefs(), "url": "https://acme.example/jobs/1",
                "location": "Remote (US)"}
        args.update(kw)
        return intake.add_pasted(self.con, **args)

    def test_verbatim(self):
        added = self.paste()
        stored, = self.con.execute("SELECT description FROM jobs WHERE id = ?",
                                   (added.job_id,)).fetchone()
        self.assertEqual(stored, self.TEXT.strip())

    def test_pasting_it_again_is_the_same_row(self):
        self.assertEqual(self.paste().job_id, self.paste().job_id)
        self.assertEqual(self.jobs(), 1)

    def test_a_title_alone_is_not_a_posting(self):
        with self.assertRaises(IntakeError):
            self.paste(text="Support Engineer, Remote")
        self.assertEqual(self.jobs(), 0)

    def test_it_needs_a_company_and_a_title(self):
        for missing in ("company", "title"):
            with self.subTest(missing=missing), self.assertRaises(IntakeError):
                self.paste(**{missing: "  "})

    def test_the_link_must_be_a_web_address(self):
        with self.assertRaises(IntakeError):
            self.paste(url="javascript:alert(1)")


class TestG_PreviewShowsTheFile(unittest.TestCase):
    """The page is the .docx's own text and formatting, escaped, and only for
    files inside output/ -- the same rule as the download."""

    def setUp(self):
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
        from tests.test_web_documents import Sandbox, make_client
        self.box = Sandbox()
        self.addCleanup(self.box.cleanup)
        doc = Document()
        name = doc.add_paragraph()
        name.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = name.add_run("Jordan Example")
        run.bold, run.font.size = True, Pt(18)
        doc.add_paragraph("<script>alert(1)</script>")
        doc.add_paragraph("Ran the support queue.", style="List Bullet")
        doc.add_paragraph().add_run("2021 - 2024").italic = True
        path = self.box.out / "resume-v1.docx"
        doc.save(path)
        self.doc_id = self.box.add_document(path)
        self.client, _ = make_client(self.box.db, self.box.out)

    def page(self):
        response = self.client.get(f"/document/{self.doc_id}/preview")
        self.assertEqual(response.status_code, 200)
        return response.text

    def test_the_words_and_their_formatting(self):
        html = self.page()
        self.assertIn("text-align:center", html)
        self.assertIn("font-size:18.0pt;font-weight:700\">Jordan Example", html)
        self.assertIn("<li><span", html)
        self.assertIn("font-style:italic\">2021 - 2024", html)

    def test_markup_in_the_document_is_shown_as_text(self):
        html = self.page()
        self.assertNotIn("<script>alert(1)</script>", html)
        self.assertIn("&lt;script&gt;", html)

    def test_a_file_outside_output_is_not_previewed(self):
        outside = self.box.dir / "resume-elsewhere.docx"
        shutil.copy(self.box.out / "resume-v1.docx", outside)
        doc_id = self.box.add_document(outside, version=2)
        self.assertEqual(self.client.get(f"/document/{doc_id}/preview").status_code, 404)

    def test_an_unknown_document_is_the_same_answer(self):
        self.assertEqual(self.client.get("/document/999/preview").status_code, 404)

    def test_the_job_page_links_to_it(self):
        self.assertIn(f"/document/{self.doc_id}/preview", self.client.get("/job/1").text)


class TestTheAddPage(unittest.TestCase):
    def setUp(self):
        from tests.test_web_documents import Sandbox, make_client
        self.box = Sandbox()
        self.addCleanup(self.box.cleanup)
        self.client, self.app = make_client(
            self.box.db, self.box.out,
            profile={"job_search_preferences": {"target_titles": ["Support Engineer"],
                                                "locations": ["Remote (US)"],
                                                "compensation_floor_usd": "no_floor"}})
        self.csrf = self.app.state.csrf_token

    def test_it_is_on_the_nav(self):
        self.assertIn('href="/add"', self.client.get("/").text)

    def test_a_refusal_comes_back_to_the_form_with_the_reason(self):
        with mock.patch("jsa.sources.fetch") as fetch:
            response = self.client.post("/add/link", data={
                "csrf": self.csrf, "url": "https://www.linkedin.com/jobs/view/1"},
                follow_redirects=False)
        fetch.assert_not_called()
        self.assertEqual(response.status_code, 303)
        self.assertIn("/add?", response.headers["location"])

    def test_a_paste_lands_on_the_new_job_page(self):
        with mock.patch.object(intake, "enrich", return_value="not checked: test"):
            response = self.client.post("/add/paste", data={
                "csrf": self.csrf, "company": "Globex", "title": "Support Engineer",
                "text": "Help customers with Python. " * 20, "location": "Remote (US)"},
                follow_redirects=False)
        self.assertEqual(response.status_code, 303)
        self.assertRegex(response.headers["location"], r"^/job/\d+\?")

    def test_the_form_needs_the_token(self):
        response = self.client.post("/add/paste", data={
            "company": "Globex", "title": "T", "text": "x" * 300})
        self.assertEqual(response.status_code, 403)


if __name__ == "__main__":
    unittest.main()
