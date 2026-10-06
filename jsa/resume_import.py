"""Import a resume (.docx or PDF) into a DRAFT profile, and preview its matches.

Rules (docs/decisions/0022-importing-a-resume.md):

- **Copy, never reword.** Every bullet, employer, title, school, project and
  certification in the draft appears word for word in the resume. `verify()`
  drops anything that doesn't, and says so. This is stricter than
  `tailor.verify_draft` on purpose: an import has nothing to shorten.
- **Never write master_profile.yaml.** The draft goes to
  `profile/master_profile.draft.yaml`; adopting it is a copy the user makes.
- **Identity never enters a prompt.** Name, email, phone, street, ZIP, links
  and the home "City, ST" are found locally and replaced with `[contact]`
  before the model call, and `tailor.scrub_prompt` checks the prompt against
  them (fail closed).
- **Degree status is never guessed.** `credential` is always left empty for
  the user to write; the resume's own education line is shown beside it.
- .docx, or a PDF with a text layer (Plan 10: `read_pdf`). A PDF that
  loses much of its text to its layout says so (`WEAK_PDF`).
"""

from __future__ import annotations

import datetime as dt
import json
import re
import sqlite3
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from . import llm
from .config import ConfigError, Preferences

REDACTED = "[contact]"
FAMILIES = ("ai_engineering", "data_ml", "sales", "technical_field")
DOCX_ONLY = ("Save your resume as .docx (Word: File → Save As) "
             "and run this again.")

# Lines that open a resume section. The contact block is everything before the
# first one; a "City, ST" or a street is only identity inside that block, since
# every job on the page has a location too.
_SECTION = re.compile(
    r"^\s*(professional\s+)?(summary|profile|objective|experience|work\s+"
    r"experience|employment|projects?|education|skills|certifications?|"
    r"technical\s+skills|licenses?)\b", re.I)

_EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
_PHONE = re.compile(
    r"(?<!\d)(?:\+?1[\s.-]?)?\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4}(?!\d)")
_URL = re.compile(
    r"(?:https?://)?(?:www\.)?(?:linkedin\.com|github\.com)/[^\s,|]+"
    r"|https?://[^\s,|]+", re.I)
# Not inside a word or link: a handle such as github.com/name12345 has five
# digits too (found on the first real run).
_ZIP = re.compile(r"(?<![\w/.@-])\d{5}(?:-\d{4})?(?![\w/])")
_CITY_ST = re.compile(r"\b([A-Z][A-Za-z.'-]+(?: [A-Z][A-Za-z.'-]+){0,2}), ([A-Z]{2})\b")
_STREET = re.compile(
    r"\b\d{1,6} [A-Za-z0-9.' ]{2,40}? (?:St|Street|Ave|Avenue|Rd|Road|Blvd|"
    r"Boulevard|Dr|Drive|Ln|Lane|Way|Ct|Court|Pl|Place|Ter|Terrace|Pkwy|"
    r"Circle|Cir)\b\.?(?:,? (?:Apt|Unit|Suite|#) ?\w+)?")
# Name particles that are lowercase inside a name: Ludwig van Beethoven.
_PARTICLES = {"van", "von", "de", "der", "den", "da", "das", "dos", "del", "della",
              "di", "du", "la", "le", "bin", "binti", "al", "el", "y", "e", "ter"}
# What a name line may carry after the name: ", MBA", " | Data Analyst".
_NAME_TAIL = re.compile(r"\s*(?:[|•·–—]|,\s|\s-\s)")
# Unzipped, a resume's .docx is well under this; more is a zip bomb (R-19).
MAX_DOCX_UNZIPPED = 50_000_000
_BULLET = re.compile(r"^\s*[•●▪◦‣∙·*–—-]\s+")
_ONGOING = {"present", "current", "now", "today", "ongoing"}
_LABEL_CONTACT = re.compile(r"\s*[-|,:–—]?\s*\[contact\]")
_YEAR = re.compile(r"(?:19|20)\d{2}")


class ResumeReadError(ConfigError):
    """The file can't be imported (wrong type, unreadable). Exit code 2."""


# --- reading ----------------------------------------------------------------


def _clean(text: str) -> str:
    text = _BULLET.sub("", text.replace(" ", " "))
    return " ".join(text.split())


def read_docx(path: Path) -> list[str]:
    """Lines from the header, body (paragraphs and tables, in order) and footer.

    Resumes often put contact details in the page header or lay the page out
    as a table; `render.extract_text` reads body paragraphs only.
    """
    path = Path(path)
    if path.suffix.lower() != ".docx":
        raise ResumeReadError(f"{path.name} is not a .docx file. {DOCX_ONLY}")
    if not path.exists():
        raise ResumeReadError(f"no file at {path}")
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        import zipfile
        with zipfile.ZipFile(path) as archive:
            unzipped = sum(i.file_size for i in archive.infolist())
    except Exception:  # noqa: BLE001 - not a zip: Document() says why below
        unzipped = 0
    if unzipped > MAX_DOCX_UNZIPPED:
        raise ResumeReadError(f"{path.name} unpacks to {unzipped // 1_000_000} MB; a "
                              "resume is far smaller. Is this the right file?")
    try:
        doc = Document(str(path))
    except Exception as exc:  # python-docx raises several unrelated types
        raise ResumeReadError(f"{path.name} could not be read as Word: {exc}. "
                           f"{DOCX_ONLY}") from exc

    lines: list[str] = []

    def add(text: str) -> None:
        for part in text.splitlines():
            part = _clean(part)
            if part and (not lines or lines[-1] != part):
                lines.append(part)

    def table(tbl) -> None:
        seen: set[int] = set()
        for row in tbl.rows:
            for cell in row.cells:
                # A merged cell is returned once per grid column it spans.
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                for p in cell.paragraphs:
                    add(p.text)
                for inner in cell.tables:
                    table(inner)

    def block(container) -> None:
        for child in container.iter_inner_content():
            if isinstance(child, Paragraph):
                add(child.text)
            elif isinstance(child, Table):
                table(child)

    for section in doc.sections:
        if not section.header.is_linked_to_previous:
            block(section.header)
    block(doc)
    for section in doc.sections:
        if not section.footer.is_linked_to_previous:
            block(section.footer)
    return lines


# PDF (Plan 10). pypdf is pure Python; it parses an untrusted file, so the
# size and page count are checked before and right after opening.
MAX_PDF_BYTES = 5_000_000
MAX_PDF_PAGES = 10            # a resume is 1-3 pages
# Page content, decompressed. The test fixtures are a few KB a page; a dense
# two-page resume is well under 100 KB. 2 MB is twenty times that.
MAX_PDF_CONTENT = 2_000_000
MIN_PDF_WORDS = 50            # under this over all pages: a scan, no text layer
SCAN = ("this PDF looks like a scan, with no text to read. Save your resume "
        "as .docx, or export a text PDF from your editor.")
# Bullet glyphs as PDF text extraction returns them: U+007F is ReportLab's
# Helvetica bullet, U+F0B7 / U+F0A7 are Word's Symbol and Wingdings bullets.
_PDF_BULLETS = "\x7f•●▪■‣⁃◦"
_PAGE_NUMBER = re.compile(r"^(?:page\s+)?\d{1,2}(?:\s*(?:of|/)\s*\d{1,2})?$", re.I)


def _pdf_line(text: str) -> str:
    """NFKC folds ligatures (U+FB01 -> "fi"); bullets and control characters
    go; whitespace is folded as for .docx."""
    import unicodedata

    text = unicodedata.normalize("NFKC", text)
    text = "".join(" " if (ch in _PDF_BULLETS or unicodedata.category(ch) == "Cc")
                   else ch for ch in text)
    return _clean(text)


def _bulletish(page_texts: list[str]) -> int:
    # Non-empty first: "" is `in` every string, and layout mode's blank lines
    # were counted as bullets (found re-measuring, 2026-10-02).
    return sum(1 for t in page_texts for line in t.splitlines()
               if line.strip() and line.lstrip()[0] in _PDF_BULLETS + "-*")


def read_pdf(path: Path) -> list[str]:
    """Lines of a text PDF. Refuses a scan, an encrypted or oversized file, or
    more than MAX_PDF_PAGES pages. No OCR.

    Extraction mode, measured 2026-10-02 (ADR 0022): on the test fixtures
    (tests/pdf_fixtures.py) pypdf's "plain" and "layout" modes give the same
    lines, and neither reconstructs two columns. On the owner's own two-page
    PDF both start 16 lines with a bullet, but plain splits text that shares a
    baseline (81 lines to layout's 65) and broke one certification name, which
    the verbatim check then dropped. So layout, unless plain finds more lines
    that start with a bullet.
    """
    path = Path(path)
    if path.stat().st_size > MAX_PDF_BYTES:
        raise ResumeReadError(f"{path.name} is over {MAX_PDF_BYTES // 1_000_000} MB; "
                              "a resume is far smaller. Is this the right file?")
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ResumeReadError(f"{path.name} is password-protected. Remove "
                                  "the password and try again.")
        pages = list(reader.pages)
    except ResumeReadError:
        raise
    except Exception as exc:  # pypdf raises many types on a damaged file
        raise ResumeReadError(f"{path.name} could not be read as a PDF ({exc}). "
                              f"{DOCX_ONLY}") from exc
    if len(pages) > MAX_PDF_PAGES:
        raise ResumeReadError(f"{path.name} has {len(pages)} pages; a resume is "
                              "1-3. Is this the right file?")
    # Text extraction costs about a second per MB of page content, twice (two
    # modes), and a few-KB file can expand to 75 MB per stream (review R-20).
    try:
        content = sum(len(p.get_contents().get_data()) for p in pages
                      if p.get_contents() is not None)
    except Exception as exc:  # noqa: BLE001 - a broken content stream
        raise ResumeReadError(f"{path.name}: its pages could not be read "
                              f"({exc}). {DOCX_ONLY}") from exc
    if content > MAX_PDF_CONTENT:
        raise ResumeReadError(f"{path.name} holds {content // 1_000_000} MB of page "
                              "content; a resume is far smaller. Is this the right file?")
    try:
        by_mode = {mode: [p.extract_text(extraction_mode=mode) or "" for p in pages]
                   for mode in ("plain", "layout")}
    except Exception as exc:  # noqa: BLE001 - a broken content stream
        raise ResumeReadError(f"{path.name}: its text could not be extracted "
                              f"({exc}). {DOCX_ONLY}") from exc
    texts = (by_mode["plain"] if _bulletish(by_mode["plain"]) > _bulletish(by_mode["layout"])
             else by_mode["layout"])

    page_lines = [[_pdf_line(raw) for raw in t.splitlines()] for t in texts]
    page_lines = [[ln for ln in lines if ln and not _PAGE_NUMBER.match(ln)]
                  for lines in page_lines]
    # A running header or footer: the same line on every page of several.
    repeated = (set.intersection(*(set(lines) for lines in page_lines))
                if len(page_lines) > 1 else set())
    lines: list[str] = []
    for n, page in enumerate(page_lines):
        lines += [ln for ln in page if n == 0 or ln not in repeated]
    if sum(len(ln.split()) for ln in lines) < MIN_PDF_WORDS:
        raise ResumeReadError(f"{path.name}: {SCAN}")
    return lines


# --- identity, found and removed locally -------------------------------------


@dataclass
class Identity:
    full_name: str | None = None
    email: str | None = None
    phone: str | None = None
    street: str | None = None
    city: str | None = None
    state: str | None = None
    postal_code: str | None = None
    links: dict[str, str | None] = field(default_factory=lambda: {
        "linkedin": None, "github": None, "portfolio": None, "website": None})

    def as_profile(self) -> dict[str, Any]:
        """The shape `tailor.identity_values` / `scrub_prompt` read."""
        return {
            "identity": {
                "full_name": self.full_name, "email": self.email,
                "phone": self.phone,
                "location": {"street": self.street, "postal_code": self.postal_code},
            },
            "links": dict(self.links),
        }


def name_on(line: str) -> str | None:
    """The name a contact-block line starts with, or None.

    2-5 words, each capitalized (any alphabet) or a lowercase particle, before
    any ", MBA" or " | headline". The ASCII-only pattern this replaces missed
    "José García", "Jane Doe, MBA" and "Ludwig van Beethoven", and the name
    then went to the model unredacted (review R-10).
    """
    head = _NAME_TAIL.split(line.strip(), maxsplit=1)[0].strip()
    words = head.split()
    if not 2 <= len(words) <= 5:
        return None
    capitals = 0
    for word in words:
        if word.lower() in _PARTICLES and word.islower():
            continue
        if not (word[0].isupper() and all(c.isalpha() or c in "'.-" for c in word)):
            return None
        capitals += 1
    return head if capitals >= 2 else None


def split_identity(lines: list[str], known: list[str] | None = None
                   ) -> tuple[Identity, list[str]]:
    """Find contact details and replace each with [contact]. Nothing is guessed:
    a field that isn't clearly there stays None and the draft marks it TODO.

    `known`: identity values already on file (the live profile's), redacted
    wherever they appear, so a name these patterns miss still never leaves."""
    ident = Identity()
    found: list[str] = [v for v in (known or []) if v]
    head = 0
    while head < len(lines) and not _SECTION.match(lines[head]):
        head += 1
    contact_block = lines[:head]

    for line in lines:
        for m in _EMAIL.finditer(line):
            ident.email = ident.email or m.group(0)
            found.append(m.group(0))
        for m in _URL.finditer(line):
            url = m.group(0).rstrip(".);")
            low = url.lower()
            key = ("linkedin" if "linkedin.com" in low else
                   "github" if "github.com" in low else "portfolio")
            if ident.links[key] is None:
                ident.links[key] = url if low.startswith("http") else "https://" + url
            found.append(url)
        for m in _PHONE.finditer(line):
            ident.phone = ident.phone or m.group(0).strip()
            found.append(m.group(0).strip())

    for line in contact_block:
        for m in _STREET.finditer(line):
            ident.street = ident.street or m.group(0)
            found.append(m.group(0))
        for m in _ZIP.finditer(line):
            ident.postal_code = ident.postal_code or m.group(0)
            found.append(m.group(0))
        m = _CITY_ST.search(line)
        if m and ident.city is None:
            ident.city, ident.state = m.group(1), m.group(2)
            found.append(m.group(0))
        name = None if _SECTION.match(line) else name_on(line)
        if ident.full_name is None and name:
            ident.full_name = name
            found.append(name)

    redacted = []
    # Longest first, so a URL containing the name is replaced whole.
    for line in lines:
        for value in sorted(set(found), key=len, reverse=True):
            line = re.sub(re.escape(value), REDACTED, line, flags=re.I)
        redacted.append(line)
    return ident, redacted


# --- the one model call -------------------------------------------------------


SYSTEM = (
    "You split a resume into structured JSON. COPY, NEVER REWORD: every "
    "string you return for company, title, location, bullets, project name, "
    "certification name/issuer, institution, field, skills and the education "
    "line must be copied character for character from the resume text. Do "
    "not shorten, merge, fix typos or improve wording. If you are not sure "
    "where something belongs, leave it out. Never write a degree or "
    "credential; that is the person's to state. Text shown as [contact] is "
    "removed contact detail: never copy it. Return only JSON."
)

SHAPE = """Return this JSON object:
{
  "experience": [{"company": "", "title": "", "location": null, "start": null,
                  "end": null, "current": false, "family": null,
                  "bullets": [{"text": "", "tags": []}]}],
  "projects": [{"name": "", "family": null, "bullets": [{"text": "", "tags": []}]}],
  "certifications": [{"name": "", "issuer": null, "issued": null}],
  "education": [{"institution": "", "field": null, "start": null, "end": null,
                 "line": ""}],
  "skills": [],
  "target_titles": []
}
- start/end: as written in the resume (e.g. "2021-03", "2021"), or null.
- bullets: one entry per bullet point, copied exactly.
- education.line: the resume's whole education line for that school, copied.
- SUGGESTIONS (these are not copied, keep them short):
  - tags: 2-6 lowercase-hyphenated keywords per bullet (e.g. "customer-facing",
    "python", "account-management").
  - family: one of %s, or null.
  - target_titles: 3-8 job titles this person could apply for.
""" % ", ".join(FAMILIES)


def extract(redacted_lines: list[str], identity: Identity,
            also: dict[str, Any] | None = None) -> dict[str, Any]:
    from .tailor import scrub_prompt

    prompt = SHAPE + "\nRESUME:\n" + "\n".join(redacted_lines)
    scrub_prompt(prompt, identity.as_profile())
    if also:                       # the live profile's identity, when one exists
        scrub_prompt(prompt, also)
    data, _usage = llm.complete_json(prompt, system=SYSTEM, max_tokens=6000,
                                     purpose="resume-import")
    if not isinstance(data, dict):
        raise llm.LLMError("the model did not return a JSON object")
    return data


# --- verification: word for word, per item, fail closed ----------------------


_GLYPHS = str.maketrans({"‘": "'", "’": "'", "“": '"',
                         "”": '"', "–": "-", "—": "-",
                         "•": " ", " ": " "})
NOT_VERBATIM = "not found word for word in your resume"


def normalize(text: str) -> str:
    return " ".join(_BULLET.sub("", str(text).translate(_GLYPHS)).lower().split())


@dataclass
class Dropped:
    where: str
    what: str
    why: str


@dataclass
class Verified:
    experience: list[dict[str, Any]] = field(default_factory=list)
    projects: list[dict[str, Any]] = field(default_factory=list)
    certifications: list[dict[str, Any]] = field(default_factory=list)
    education: list[dict[str, Any]] = field(default_factory=list)
    skills: list[str] = field(default_factory=list)
    target_titles: list[str] = field(default_factory=list)
    dropped: list[Dropped] = field(default_factory=list)

    @property
    def bullet_count(self) -> int:
        return sum(len(e["bullets"]) for e in self.experience + self.projects)

    @property
    def unmatched_share(self) -> float:
        """Of the bullets the model found, the share not in the text verbatim."""
        missed = sum(1 for d in self.dropped
                     if d.where.endswith("bullet") and d.why == NOT_VERBATIM)
        found = self.bullet_count + missed
        return missed / found if found else 0.0


def _str(value: Any) -> str:
    return value.strip() if isinstance(value, str) else ""


def _tags(value: Any) -> list[str]:
    out = []
    for t in value if isinstance(value, list) else []:
        t = re.sub(r"[^a-z0-9+#.-]+", "-", _str(t).lower()).strip("-")
        if t and t not in out:
            out.append(t)
    return out[:8]


def verify(extracted: dict[str, Any], resume_text: str) -> Verified:
    """Keep only what the resume says word for word. Tags, family and target
    titles are suggestions and are marked so in the draft, not checked."""
    # A PDF breaks a long word across lines with a hyphen ("coordin-" /
    # "ated"). Two variants undo only that line break, adding no text: one
    # rejoins the word, one keeps a real hyphen ("40-" / "person").
    hays = (normalize(resume_text),
            normalize(re.sub(r"(?<=[a-z])-\n(?=[a-z])", "", resume_text)),
            normalize(re.sub(r"-\n", "-", resume_text)))
    years = set(_YEAR.findall(resume_text))
    out = Verified()

    def ok(text: str) -> bool:
        if not text or REDACTED in text:
            return False
        needle = normalize(text)
        return any(needle in hay for hay in hays)

    def keep(where: str, value: Any, label: bool = False) -> str | None:
        text = _str(value)
        if label:
            # A heading such as "Project Name - <repo link>": the link was
            # removed as contact detail, and the name alone is what counts.
            text = _LABEL_CONTACT.sub("", text).strip(" -|,:–—")
        if not text:
            return None
        if ok(text):
            return text
        out.dropped.append(Dropped(where, text, NOT_VERBATIM))
        return None

    def date(where: str, value: Any) -> str | None:
        text = _str(str(value)) if value is not None else ""
        if not text or text.lower() in _ONGOING:
            return None  # "Present" is not a date; `current` carries it
        found = _YEAR.findall(text)
        if found and all(y in years for y in found):
            return text
        out.dropped.append(Dropped(where, text, "no year in it appears in your resume"))
        return None

    def bullets(where: str, items: Any) -> list[dict[str, Any]]:
        kept = []
        for b in items if isinstance(items, list) else []:
            text = _str(b.get("text") if isinstance(b, dict) else b)
            if not text:
                continue
            if ok(text):
                kept.append({"text": text, "tags": _tags(
                    b.get("tags") if isinstance(b, dict) else [])})
            else:
                out.dropped.append(Dropped(where + " bullet", text, NOT_VERBATIM))
        return kept

    def lost(entry: str, items: Any, why: str) -> None:
        for b in items if isinstance(items, list) else []:
            text = _str(b.get("text") if isinstance(b, dict) else b)
            if text:
                out.dropped.append(Dropped(entry + " bullet", text, why))

    def family(value: Any) -> str | None:
        return value if value in FAMILIES else None

    mentions_present = bool(re.search(r"\b(present|current|now)\b", hays[0]))
    for e in _items(extracted.get("experience")):
        company = keep("employer", e.get("company"), label=True)
        title = keep("job title", e.get("title"), label=True)
        if not company or not title:
            if company or title:
                out.dropped.append(Dropped(
                    "job", company or title,
                    "the job's employer and title must both be in your resume"))
            lost(company or title or "job", e.get("bullets"),
                 "dropped with its job")
            continue
        out.experience.append({
            "company": company, "title": title,
            "location": keep("job location", e.get("location")),
            "start": date(f"{company} start", e.get("start")),
            "end": date(f"{company} end", e.get("end")),
            "current": (bool(e.get("current"))
                        or _str(e.get("end")).lower() in _ONGOING) and mentions_present,
            "family": family(e.get("family")),
            "bullets": bullets(company, e.get("bullets")),
        })
    for p in _items(extracted.get("projects")):
        name = keep("project", p.get("name"), label=True)
        if name:
            out.projects.append({"name": name, "family": family(p.get("family")),
                                 "bullets": bullets(name, p.get("bullets"))})
        else:
            lost("project", p.get("bullets"), "dropped with its project")
    for c in _items(extracted.get("certifications")):
        name = keep("certification", c.get("name"), label=True)
        if name:
            out.certifications.append({
                "name": name, "issuer": keep("issuer", c.get("issuer")),
                "issued": date(f"{name} issued", c.get("issued"))})
    for ed in _items(extracted.get("education")):
        school = keep("school", ed.get("institution"), label=True)
        if school:
            # Any degree the model offers is ignored, never checked: the
            # credential is the person's to write (check_credentials).
            out.education.append({
                "institution": school, "field": keep("field of study", ed.get("field")),
                "start": date(f"{school} start", ed.get("start")),
                "end": date(f"{school} end", ed.get("end")),
                "line": keep("education line", ed.get("line"))})
    for s in extracted.get("skills") or []:
        s = keep("skill", s)
        if s and s not in out.skills:
            out.skills.append(s)
    for t in extracted.get("target_titles") or []:
        t = _str(t)
        if t and REDACTED not in t and len(t) <= 60 and t not in out.target_titles:
            out.target_titles.append(t)
    return out


def _items(value: Any) -> list[dict[str, Any]]:
    return [v for v in value if isinstance(v, dict)] if isinstance(value, list) else []


# --- the draft file, written by hand to keep its comments -------------------


SUGGESTED = "  # suggested: check"


def _q(value: Any) -> str:
    """A YAML scalar. JSON strings are valid YAML double-quoted scalars."""
    if value is None:
        return "null"
    if isinstance(value, bool):
        return "true" if value else "false"
    return json.dumps(str(value), ensure_ascii=False)


def _slug(text: str, used: set[str]) -> str:
    base = "_".join(re.findall(r"[a-z0-9]+", text.lower())[:2])[:20] or "entry"
    slug, n = base, 2
    while slug in used:
        slug, n = f"{base}{n}", n + 1
    used.add(slug)
    return slug


def render_draft(ident: Identity, v: Verified, source: str) -> str:
    L: list[str] = []
    w = L.append
    w("# DRAFT profile imported from your resume by `jsa import-resume`.")
    w("# Nothing here is live until YOU copy it to profile/master_profile.yaml.")
    w("#")
    w("# Every bullet, employer, title and school below was found word for word")
    w("# in your resume. Lines marked `# suggested: check` were proposed by the")
    w("# model (tags, families, target titles): keep, change or delete them.")
    w("# Lines marked TODO are decisions only you can make; `jsa doctor` lists them.")
    w("# This file is gitignored: it holds your contact details.")
    w("")
    w("schema_version: 1")
    w(f"last_updated: {dt.date.today().isoformat()}")
    w(f"source_document: {_q(source)}")
    w("")
    w("identity:")
    w(f"  full_name: {_q(ident.full_name)}" + ("" if ident.full_name else "  # TODO"))
    w("  preferred_name: null")
    w("  pronouns: null")
    w("  location:")
    w(f"    street: {_q(ident.street)}")
    w(f"    city: {_q(ident.city)}")
    w(f"    state: {_q(ident.state)}")
    w(f"    postal_code: {_q(ident.postal_code)}")
    w("    country: null")
    w(f"  email: {_q(ident.email)}" + ("" if ident.email else "  # TODO"))
    w(f"  phone: {_q(ident.phone)}" + ("" if ident.phone else "  # TODO"))
    w("")
    w("links:")
    for key in ("linkedin", "github", "portfolio", "website"):
        w(f"  {key}: {_q(ident.links.get(key))}")
    w("")
    w("job_search_preferences:")
    w("  target_titles:" + ("" if v.target_titles else " []  # TODO"))
    for t in v.target_titles:
        w(f"    - {_q(t)}{SUGGESTED}")
    w("  fallback_titles: []")
    w("  fallback_weight: 0.7")
    w(f"  seniority: [\"intern\", \"entry\", \"junior\", \"associate\", \"mid\"]{SUGGESTED}")
    w(f"  work_arrangement: [\"remote\", \"hybrid\", \"onsite\"]{SUGGESTED}")
    w("  locations:")
    w(f"    - \"Remote (US)\"{SUGGESTED}")
    if ident.city and ident.state:
        w(f"    - {_q(f'{ident.city}, {ident.state}')}{SUGGESTED}")
    w("  compensation_floor_usd: null  # TODO no_floor | a number (see the example file)")
    w("  work_authorization: null      # TODO e.g. \"US citizen - no sponsorship required\"")
    w("  needs_visa_sponsorship: null  # TODO true/false")
    w("  willing_to_relocate: null     # TODO true/false")
    w("  max_years_experience: 3       # TODO the years you can credibly claim")
    w("  years_filter: reject")
    w("  exclude_keywords: []          # TODO words that rule a posting out, e.g. \"Senior\"")
    w("")
    w("summaries: []  # TODO at least one: it opens every document. Write it yourself.")
    w("")

    used: set[str] = set()

    def entry_bullets(slug: str, items: list[dict[str, Any]]) -> None:
        if not items:
            w("    bullets: []  # TODO none were found word for word")
            return
        w("    bullets:")
        for n, b in enumerate(items, 1):
            w(f"      - id: b_{slug}_{n}")
            w("        strength: 1  # TODO rate 1-3 (1 = always include)")
            w(f"        text: {_q(b['text'])}")
            w(f"        tags: [{', '.join(_q(t) for t in b['tags'])}]{SUGGESTED}")

    def fam(value: str | None) -> str:
        return (f"{value}{SUGGESTED}" if value else
                f"null  # TODO one of {', '.join(FAMILIES)}")

    w("experience:" + ("" if v.experience else " []"))
    for e in v.experience:
        slug = _slug(e["company"], used)
        w(f"  - id: exp_{slug}")
        w(f"    company: {_q(e['company'])}")
        w(f"    title: {_q(e['title'])}")
        w(f"    location: {_q(e['location'])}")
        w(f"    start: {_q(e['start'])}")
        w(f"    end: {_q(e['end'])}")
        w(f"    current: {_q(e['current'])}")
        w(f"    family: {fam(e['family'])}")
        entry_bullets(slug, e["bullets"])
    w("")
    w("projects:" + ("" if v.projects else " []"))
    for p in v.projects:
        slug = _slug(p["name"], used)
        w(f"  - id: proj_{slug}")
        w(f"    name: {_q(p['name'])}")
        w("    status: null  # TODO released | in_development")
        w("    repo: null")
        w(f"    family: {fam(p['family'])}")
        entry_bullets(slug, p["bullets"])
    w("")
    w("certifications:" + ("" if v.certifications else " []"))
    for n, c in enumerate(v.certifications, 1):
        w(f"  - id: cert_{n}")
        w(f"    name: {_q(c['name'])}")
        w(f"    issuer: {_q(c['issuer'])}")
        w(f"    issued: {_q(c['issued'])}")
    w("")
    w("education:" + ("" if v.education else " []"))
    for n, ed in enumerate(v.education, 1):
        w(f"  - id: edu_{n}")
        w(f"    institution: {_q(ed['institution'])}")
        if ed["line"]:
            w(f"    # your resume says: {ed['line']}")
        w("    credential: null  # TODO write exactly what you hold; "
          "nothing guesses this")
        w(f"    field: {_q(ed['field'])}")
        w(f"    start: {_q(ed['start'])}")
        w(f"    end: {_q(ed['end'])}")
    w("")
    w("skills:")
    w("  # From your resume. Sort them into ai_tools / technical / applied_focus.")
    w("  technical:" + ("" if v.skills else " []"))
    for s in v.skills:
        w(f"    - {_q(s)}")
    w("")
    w("ats_keywords:")
    w("  have:" + ("" if v.skills else " []"))
    for s in v.skills:
        w(f"    - {_q(s)}{SUGGESTED}")
    w("  aspirational_do_not_claim: []  # TODO skills a document must never claim")
    w("")
    w("gaps_and_notes: []")
    return "\n".join(L) + "\n"


def write_draft(path: Path, ident: Identity, v: Verified, source: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(render_draft(ident, v, source), encoding="utf-8")


# --- preview: what the draft would match, read-only ---------------------------


@dataclass
class Match:
    score: float
    job_id: int
    company: str
    title: str


def preview(con: sqlite3.Connection, draft: dict[str, Any],
            limit: int = 10) -> list[Match]:
    """Score stored open jobs against the draft, in memory. Writes nothing.

    Raises ConfigError when the draft can't be scored yet (no target titles
    or locations), which the caller reports instead of a preview.
    """
    from .scoring import score_job

    prefs = Preferences.from_profile(draft)
    rows = con.execute(
        "SELECT j.id, j.title, j.description, j.location, j.remote, "
        "j.salary_min, j.salary_max, j.salary_period, j.salary_text, "
        "j.salary_currency, j.salary_source, c.name AS company "
        "FROM jobs j JOIN companies c ON c.id = j.company_id "
        "WHERE j.archived_at IS NULL AND j.closed_at IS NULL").fetchall()
    found = []
    for row in rows:
        score, _ = score_job(dict(row), prefs)
        if score > 0:
            found.append(Match(score, row["id"], row["company"], row["title"]))
    found.sort(key=lambda m: (-m.score, m.job_id))
    return found[:limit]


# --- one entry point for the CLI and the dashboard ----------------------------


# Extension -> the bytes such a file starts with. The dashboard checks both
# and never trusts the browser's content type; the CLI checks them too.
SUPPORTED: dict[str, bytes] = {".docx": b"PK\x03\x04", ".pdf": b"%PDF-"}
KINDS = " or ".join(sorted(SUPPORTED))
DRAFT_NAME = "master_profile.draft.yaml"


class DraftRefused(ConfigError):
    """The draft can't be written where asked (the live profile, or an
    existing draft without force)."""


@dataclass
class ImportReport:
    draft_path: Path
    verified: Verified
    blocking: list[Any]          # doctor findings that block, on the draft
    matches: list[Match]
    no_preview: str = ""         # why there is no preview, when there isn't
    warning: str = ""            # the PDF lost much of its text


# Share of the bullets the model found that were not in the text word for
# word, above which a PDF import says its layout likely lost text. Measured
# 2026-10-02 on the fixtures in tests/pdf_fixtures.py: the one-column and
# ASCII85 PDFs lose 0 of 3; the two-column one loses 2 of 3 (67%). The owner's
# own PDF lost 0 of 10. Set low enough that losing one bullet in three warns.
WEAK_PDF_SHARE = 0.30
WEAK_PDF = ("much of this PDF's text could not be matched word for word, likely "
            "its layout (columns or tables). The .docx version usually imports "
            "more completely.")


def check_type(path: Path) -> None:
    """Refuse a file whose name and first bytes don't agree on a known type."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise ResumeReadError(f"{path.name} is not a {KINDS} file. {DOCX_ONLY}")
    if not path.exists():
        raise ResumeReadError(f"no file at {path}")
    with path.open("rb") as fh:
        head = fh.read(8)
    if not head.startswith(SUPPORTED[suffix]):
        raise ResumeReadError(
            f"{path.name} is named {suffix} but is not one inside. {DOCX_ONLY}")


def read(path: Path) -> list[str]:
    check_type(path)
    return read_pdf(path) if Path(path).suffix.lower() == ".pdf" else read_docx(path)


def _live_identity(live: Path) -> dict[str, Any] | None:
    """The identity and links of the profile on file, or None."""
    try:
        import yaml
        data = yaml.safe_load(Path(live).read_text(encoding="utf-8")) or {}
    except Exception:  # noqa: BLE001 - no profile, or not one we can read
        return None
    if not isinstance(data, dict):
        return None
    return {"identity": data.get("identity") or {}, "links": data.get("links") or {}}


def run(path: Path, *, out: Path, live: Path, force: bool = False,
        con: sqlite3.Connection | None = None,
        source_name: str | None = None) -> ImportReport:
    """Resume -> draft profile -> doctor's blocking items and a preview.

    Never writes `live`. Refuses before the model call, so a refusal costs
    nothing.
    """
    from . import doctor

    out, live = Path(out), Path(live)
    if out.name == live.name or out.resolve() == live.resolve():
        raise DraftRefused(f"{out} is your live profile. The import only ever "
                           "writes a draft; copying it over is yours to do.")
    if out.exists() and not force:
        raise DraftRefused(f"{out} already exists. Review it, or replace the "
                           "draft (--force, or tick 'Replace my draft').")

    lines = read(path)
    on_file = _live_identity(live)
    from .tailor import identity_values
    ident, redacted = split_identity(lines, identity_values(on_file) if on_file else None)
    verified = verify(extract(redacted, ident, on_file), "\n".join(redacted))
    write_draft(out, ident, verified, source_name or Path(path).name)

    import yaml

    draft = yaml.safe_load(out.read_text(encoding="utf-8"))
    report = ImportReport(out, verified, doctor.run(draft, con).blocking, [])
    if Path(path).suffix.lower() == ".pdf" and verified.unmatched_share > WEAK_PDF_SHARE:
        report.warning = WEAK_PDF
    if con is None:
        report.no_preview = "no tracker yet: run `jsa init` and `jsa discover` first."
    else:
        try:
            report.matches = preview(con, draft)
        except ConfigError as exc:
            report.no_preview = str(exc)
    return report
