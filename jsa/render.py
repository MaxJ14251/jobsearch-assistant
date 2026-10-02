"""Render a verified draft to .docx and record its provenance.

ATS-safe formatting is a functional requirement, not a style preference. An
applicant tracking system parses the file; anything it cannot read costs the
application. So: one column, no text boxes, no tables for layout, no content in
headers or footers, no images, real list paragraphs, parseable dates.

The source resume for this project was a ReportLab PDF whose text extraction
was genuinely difficult — ASCII85 over Flate, no text layer tooling available.
.docx exists so that never happens again, and `tests/test_render.py` proves the
text comes back out.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import json
import re
import sqlite3
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from . import db
from .config import OUTPUT_DIR, load_coach
from .tailor import TailoredDraft, collect_bullets, entry_key, public_repo

BODY_PT = 10.5
NAME_PT = 18
HEADING_PT = 11


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", (text or "").lower()).strip("-")[:60] or "item"


def contact_line(profile: dict[str, Any]) -> str:
    """Identity, merged locally. None of this was ever sent to a model."""
    ident = profile.get("identity") or {}
    loc = ident.get("location") or {}
    links = profile.get("links") or {}
    city = ", ".join(p for p in (loc.get("city"), loc.get("state")) if p)
    parts = [city, ident.get("phone"), ident.get("email")]
    for url in (links.get("linkedin"), links.get("github")):
        if url:
            parts.append(re.sub(r"^https?://", "", url))
    return "  |  ".join(p for p in parts if p)


def _heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(HEADING_PT)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)


def _body(doc: Document, text: str, *, bullet: bool = False) -> None:
    # A real list paragraph, not a literal "•" character an ATS must guess at.
    p = doc.add_paragraph(style="List Bullet" if bullet else None)
    run = p.add_run(text)
    run.font.size = Pt(BODY_PT)
    p.paragraph_format.space_after = Pt(2)


class RenderError(RuntimeError):
    """The document would misplace a bullet. Nothing is saved."""


# A fixed table, not strftime("%b"): that follows the OS locale, and a
# resume's dates must not change language with the machine that drafts it.
MONTHS = ("Jan", "Feb", "Mar", "Apr", "May", "Jun",
          "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")
DATE_STYLES = ("month_year", "numeric", "iso")


def format_date(value: Any, style: str = "month_year") -> str:
    """2022-03 -> "Mar 2022" (or "03/2022", "2022-03"); a year stays a year;
    anything else is printed as written."""
    if value is None or value == "":
        return ""
    if isinstance(value, dt.date):
        year, month = value.year, value.month
    else:
        text = str(value).strip()
        m = re.fullmatch(r"(\d{4})-(\d{1,2})(?:-\d{1,2})?", text)
        if not m or not 1 <= int(m[2]) <= 12:
            return text
        year, month = int(m[1]), int(m[2])
    if style == "numeric":
        return f"{month:02d}/{year}"
    if style == "iso":
        return f"{year}-{month:02d}"
    return f"{MONTHS[month - 1]} {year}"


def date_style(profile: dict[str, Any]) -> str:
    style = (profile.get("resume") or {}).get("date_style") or "month_year"
    if style not in DATE_STYLES:
        raise RenderError(f"resume.date_style is {style!r}; use one of "
                          f"{', '.join(DATE_STYLES)}")
    return style


def _dates(start: Any, end: Any, current: bool = False,
           style: str = "month_year") -> str:
    a = format_date(start, style)
    b = "Present" if current else format_date(end, style)
    return f"{a} – {b}".strip(" –")


def skill_label(key: str, labels: dict[str, Any], acronyms: list[str]) -> str:
    """`skill_labels[key]` if set, else the key humanized, with the acronyms
    from config/coach.yaml in capitals ("ai_tools" -> "AI Tools")."""
    if labels.get(key):
        return str(labels[key])
    upper = {str(a).lower(): str(a) for a in acronyms}
    return " ".join(upper.get(w.lower(), w.capitalize())
                    for w in re.split(r"[_\s]+", str(key)) if w)


# Skill categories the profile keeps for its owner, never for a reader.
UNPRINTED_SKILLS = ("unverified_candidates",)


@dataclass
class RenderedDoc:
    kind: str
    path: Path
    document_id: int | None = None


def render_resume(
    draft: TailoredDraft, profile: dict[str, Any], job: dict[str, Any], out: Path
) -> Path:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(BODY_PT)

    ident = profile.get("identity") or {}
    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.add_run(ident.get("full_name", ""))
    run.bold = True
    run.font.size = Pt(NAME_PT)
    name.paragraph_format.space_after = Pt(1)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    crun = contact.add_run(contact_line(profile))
    crun.font.size = Pt(9)

    style = date_style(profile)
    if draft.summary:
        _heading(doc, "Summary")
        _body(doc, draft.summary)

    # Bullets are grouped by the same key collect_bullets gives them
    # (tailor.entry_key), so the two can't disagree. They did: this used
    # `id or company` while bullets were keyed `company :: title` when an
    # entry had no id, and the whole Experience section silently vanished.
    sources = collect_bullets(profile)
    by_parent: dict[str, list[tuple[str, str]]] = {}
    for b in draft.bullets:
        src = sources.get(b.source_id)
        if src:
            by_parent.setdefault(f"{src.origin}::{src.parent}", []).append(
                (b.source_id, b.text))
    placed: list[tuple[str, str]] = []   # (bullet id, entry it was written under)

    def write_bullets(where: str, items: list[tuple[str, str]]) -> None:
        for source_id, text in items:
            _body(doc, text, bullet=True)
            placed.append((source_id, where))

    def write_experience() -> None:
        written = False
        for exp in profile.get("experience") or []:
            where = f"experience::{entry_key(exp, 'company', 'title')}"
            texts = by_parent.get(where)
            if not texts:
                continue
            if not written:
                _heading(doc, "Experience")
                written = True
            line = doc.add_paragraph()
            r = line.add_run(f"{exp.get('title')} — {exp.get('company')}")
            r.bold = True
            r.font.size = Pt(BODY_PT)
            meta = doc.add_paragraph()
            m = meta.add_run(
                "  |  ".join(
                    p for p in (exp.get("location"), exp.get("company_descriptor"),
                                _dates(exp.get("start"), exp.get("end"),
                                       exp.get("current", False), style)) if p
                )
            )
            m.italic = True
            m.font.size = Pt(9)
            meta.paragraph_format.space_after = Pt(1)
            write_bullets(where, texts)

    def write_projects() -> None:
        written = False
        for proj in profile.get("projects") or []:
            where = f"project::{entry_key(proj, 'name', 'title')}"
            texts = by_parent.get(where)
            if not texts:
                continue
            if not written:
                _heading(doc, "Projects")
                written = True
            line = doc.add_paragraph()
            label = proj.get("name", "")
            # Status is shown, never quietly dropped to make a project look shipped.
            if proj.get("status") == "in_development":
                label += " (in development)"
            r = line.add_run(label)
            r.bold = True
            r.font.size = Pt(BODY_PT)
            # The project's public link, as plain text: a reader can type it,
            # and an ATS reads it, without a hyperlink field either has to parse.
            repo = public_repo(proj)
            if repo:
                link = line.add_run(" — " + repo.split("://", 1)[1].rstrip("/"))
                link.font.size = Pt(BODY_PT)
            line.paragraph_format.space_after = Pt(1)
            write_bullets(where, texts)

    # Whichever section holds more of the selected bullets leads. Experience
    # used to lead unconditionally, so a resume for a robotics role opened with
    # fire-alarm installation while the relevant AI projects sat below it.
    # Experience wins ties, because for most roles and most readers that is the
    # conventional and expected order.
    counts = {"experience": 0, "project": 0}
    for key, texts in by_parent.items():
        origin = key.split("::", 1)[0]
        if origin in counts:
            counts[origin] += len(texts)
    if counts["project"] > counts["experience"]:
        write_projects()
        write_experience()
    else:
        write_experience()
        write_projects()

    # Every category, in the profile's order. Only ai_tools, technical and
    # applied_focus used to print ("Ai Tools"), so any other field's skills
    # were silently dropped.
    skills = {k: v for k, v in (profile.get("skills") or {}).items()
              if k not in UNPRINTED_SKILLS and isinstance(v, list) and v}
    if skills:
        _heading(doc, "Skills")
        labels = profile.get("skill_labels") or {}
        acronyms = load_coach().get("acronyms") or []
        for key, vals in skills.items():
            _body(doc, f"{skill_label(key, labels, acronyms)}: "
                       f"{', '.join(str(v) for v in vals)}")

    certs = profile.get("certifications") or []
    if certs:
        _heading(doc, "Certifications")
        for c in certs:
            # display_name, when set, is how the operator wants a coded name
            # read. `description` stays off the page (too long for one line).
            issued = format_date(c.get("issued"), style)
            line = f"{c.get('display_name') or c.get('name')} — {c.get('issuer')}"
            line += f" ({issued})" if issued else ""
            parts = [str(x.get("name")) for x in c.get("components") or []
                     if isinstance(x, dict) and x.get("name")]
            if parts:
                line += ": " + "; ".join(parts)
            _body(doc, line)

    edu = profile.get("education") or []
    if edu:
        _heading(doc, "Education")
        for e in edu:
            field = e.get("field")
            years = _dates(e.get("start"), e.get("end"), style=style)
            # The credential string is reproduced verbatim. It says the degree
            # was not conferred, and nothing here may soften that.
            line = f"{e.get('institution')}"
            if field:
                line += f" — {field}"
            if years:
                line += f", {years}"
            credential = " ".join(str(e.get("credential") or "").split())
            if credential:
                line += f" · {credential}"
            _body(doc, line)

    check_placement(draft, sources, placed)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out


def check_placement(draft: TailoredDraft, sources: dict[str, Any],
                    placed: list[tuple[str, str]]) -> None:
    """Every drafted bullet printed exactly once, under its own entry.

    ADR 0011's "a guard for it does not exist": a bullet under the wrong
    heading is a resume claiming one job's work for another. Raises before
    anything is saved.
    """
    times = Counter(source_id for source_id, _ in placed)
    under = dict(placed)
    for b in draft.bullets:
        src = sources.get(b.source_id)
        expected = f"{src.origin}::{src.parent}" if src else None
        if times[b.source_id] != 1 or under.get(b.source_id) != expected:
            raise RenderError(
                f"bullet {b.source_id!r} was written {times[b.source_id]} time(s)"
                + (f", under {under[b.source_id]!r}" if b.source_id in under else "")
                + f"; it belongs once under {expected!r}. Nothing was saved. "
                "Two entries with the same employer and title need an id:.")


def render_cover_letter(
    draft: TailoredDraft, profile: dict[str, Any], job: dict[str, Any],
    body: str, out: Path
) -> Path:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(BODY_PT)

    ident = profile.get("identity") or {}
    head = doc.add_paragraph()
    r = head.add_run(ident.get("full_name", ""))
    r.bold = True
    r.font.size = Pt(13)
    c = doc.add_paragraph().add_run(contact_line(profile))
    c.font.size = Pt(9)

    doc.add_paragraph()
    _body(doc, f"Re: {job.get('title')} at {job.get('company', '')}".strip())
    doc.add_paragraph()
    for para in [p for p in body.split("\n\n") if p.strip()]:
        _body(doc, " ".join(para.split()))
        doc.add_paragraph()
    _body(doc, ident.get("full_name", ""))

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out


def extract_text(path: Path) -> str:
    """Pull the text back out — the same thing an ATS has to do.

    If this returns nothing useful, the document is unreadable to a parser no
    matter how good it looks on screen.
    """
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def record(
    con: sqlite3.Connection, *, job_id: int | None, kind: str, path: Path,
    draft: TailoredDraft, prompt_hash: str, note: str | None = None,
) -> int:
    """Insert a documents row with full provenance."""
    version = con.execute(
        "SELECT COALESCE(MAX(version), 0) + 1 FROM documents "
        "WHERE job_id IS ? AND kind = ?", (job_id, kind),
    ).fetchone()[0]
    cur = con.execute(
        """INSERT INTO documents
             (job_id, kind, path, format, version, bullet_ids,
              keywords_matched, keywords_missing, model, prompt_hash, note)
           VALUES (:job_id,:kind,:path,'docx',:version,:bullet_ids,
                   :matched,:missing,:model,:prompt_hash,:note)""",
        {
            "job_id": job_id, "kind": kind,
            "path": str(path), "version": version,
            "bullet_ids": json.dumps([b.source_id for b in draft.bullets]),
            "matched": json.dumps(draft.keywords_matched),
            "missing": json.dumps(draft.keywords_missing),
            "model": draft.model, "prompt_hash": prompt_hash,
            "note": note or None,
        },
    )
    return int(cur.lastrowid)


def prompt_hash(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()[:16]


def next_version(con: sqlite3.Connection, job_id: int | None, kind: str) -> int:
    """The version a new document of this kind for this job would take."""
    return int(con.execute(
        "SELECT COALESCE(MAX(version), 0) + 1 FROM documents "
        "WHERE job_id IS ? AND kind = ?", (job_id, kind),
    ).fetchone()[0])


def output_path(company: str, title: str, kind: str, version: int = 1) -> Path:
    """Where a rendered document lives.

    The version is in the FILENAME, not just the documents row. Without it a
    redraft overwrites the file its predecessor's row still points at -- and
    since approval is per version (ADR 0003 decision 3), an approved document
    would silently come to mean different content than the human approved.
    """
    name = f"{kind}.docx" if version <= 1 else f"{kind}-v{version}.docx"
    return OUTPUT_DIR / slug(company) / slug(title) / name
